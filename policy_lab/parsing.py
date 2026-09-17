"""Format-aware parsing with native locators and a killable Docling subprocess."""
from __future__ import annotations

import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import uuid
import zipfile

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm", ".md", ".txt"}


def validate_blocks(blocks, settings):
    if not isinstance(blocks, list) or not blocks:
        raise ValueError("Document contains no readable text")
    result, seen, total = [], set(), 0
    for ordinal, block in enumerate(blocks):
        if not isinstance(block, dict) or not isinstance(block.get("text"), str):
            raise ValueError("Invalid parsed block")
        text = block["text"].strip()
        if not text:
            continue
        total += len(text)
        if total > settings.max_characters:
            raise ValueError("Extracted text exceeds configured character limit")
        identifier = block.get("id") or uuid.uuid4().hex
        if identifier in seen:
            raise ValueError("Duplicate block identifier")
        seen.add(identifier)
        locator = block.get("locator", {})
        if not isinstance(locator, dict):
            raise ValueError("Invalid source locator")
        result.append(dict(id=identifier, ordinal=ordinal, heading=str(block.get("heading", "")),
                           text=text, kind=str(block.get("kind", "paragraph")), locator=locator))
    if not result:
        raise ValueError("Document contains no readable text")
    return result


def _text_blocks(text, markdown=False):
    """Paragraph/heading units retain line references; long units become bounded blocks."""
    heading = ""
    headings = {}
    current, start = [], 1
    blocks = []
    def flush(end):
        nonlocal current
        if current:
            content = "\n".join(current).strip()
            if content:
                # Bound complete evidence size; locator flags long paragraph subdivisions.
                for offset in range(0, len(content), 12000):
                    blocks.append(dict(text=content[offset:offset+12000], heading=heading,
                        kind="paragraph", locator={"line_start": start, "line_end": end,
                        "character_offset": offset, "format": "markdown" if markdown else "text"}))
            current = []
    for line_no, line in enumerate(text.splitlines(), 1):
        match = re.match(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$", line) if markdown else None
        if match:
            flush(line_no - 1)
            level = len(match.group(1))
            headings = {key: value for key, value in headings.items() if key < level}
            headings[level] = match.group(2)
            heading = " / ".join(headings[key] for key in sorted(headings))
            blocks.append(dict(text=line.strip(), heading=heading, kind="heading",
                               locator={"line_start": line_no, "line_end": line_no, "level": len(match.group(1)), "format": "markdown"}))
        elif not line.strip():
            flush(line_no - 1)
        else:
            if not current:
                start = line_no
            current.append(line)
    flush(len(text.splitlines()))
    return blocks


class DocumentParser:
    def __init__(self, settings):
        self.settings = settings

    def parse(self, path, name=None):
        path = Path(path)
        extension = Path(name or path.name).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise ValueError("Unsupported document format")
        if path.stat().st_size > self.settings.max_upload_bytes:
            raise ValueError("File exceeds configured upload limit")
        if extension in {".docx", ".pptx", ".xlsx"}:
            try:
                with zipfile.ZipFile(path) as archive:
                    entries = archive.infolist()
                    if len(entries) > 20000 or sum(entry.file_size for entry in entries) > self.settings.max_upload_bytes * 4:
                        raise ValueError("Office document exceeds expanded archive resource limit")
                    if extension == ".pptx":
                        pages = sum(bool(re.fullmatch(r"ppt/slides/slide\d+\.xml", entry.filename)) for entry in entries)
                        if pages > self.settings.max_pages:
                            raise ValueError("Document exceeds configured slide limit")
            except zipfile.BadZipFile:
                raise ValueError("Invalid Office document container") from None
        if extension in {".txt", ".md"}:
            # UTF-8 and UTF-16 BOMs only: ambiguous legacy encoding must not silently corrupt policy text.
            raw = path.read_bytes()
            encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
            try:
                text = raw.decode(encoding)
            except UnicodeDecodeError:
                raise ValueError("Text files must use UTF-8 or UTF-16 with a BOM") from None
            if "\x00" in text:
                raise ValueError("Binary content is not a text document")
            if len(text) > self.settings.max_characters:
                raise ValueError("Extracted text exceeds configured character limit")
            return validate_blocks(_text_blocks(text, extension == ".md"), self.settings)
        with tempfile.TemporaryDirectory(prefix="docling-result-", dir=self.settings.data_dir) as directory:
            output = Path(directory) / "parsed.json"
            config = {"max_pages": self.settings.max_pages, "max_characters": self.settings.max_characters,
                      "max_upload_bytes": self.settings.max_upload_bytes, "parser_timeout": self.settings.parser_timeout,
                      "model_cache_dir": str(self.settings.model_cache_dir.resolve())}
            try:
                result = subprocess.run([sys.executable, "-m", "policy_lab.parsing", str(path.resolve()),
                    str(output.resolve()), json.dumps(config)], capture_output=True, text=True,
                    encoding="utf-8", errors="replace", timeout=self.settings.parser_timeout, check=False)
            except subprocess.TimeoutExpired:
                raise ValueError("Document parsing exceeded configured timeout") from None
            if result.returncode != 0:
                # The dedicated worker emits one bounded error; never return complete subprocess traces.
                error = result.stderr.strip().splitlines()
                raise ValueError("Document parser failed: " + (error[-1][:500] if error else "No parser output"))
            return validate_blocks(json.loads(output.read_text(encoding="utf-8")), self.settings)


def _docling_convert(path, config):
    from docling.document_converter import DocumentConverter, PdfFormatOption, HTMLFormatOption
    from docling.datamodel.backend_options import HTMLBackendOptions
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions, TesseractCliOcrOptions
    extension = path.suffix.lower()
    options = {}
    if extension in {".html", ".htm"}:
        options[InputFormat.HTML] = HTMLFormatOption(backend_options=HTMLBackendOptions(
            enable_remote_fetch=False, enable_local_fetch=False, fetch_images=False, render_page=False))
    if extension == ".pdf":
        artifacts = Path(config["model_cache_dir"]) / "docling"
        if not artifacts.is_dir() or not any(artifacts.iterdir()):
            raise ValueError("Docling models are not provisioned in MODEL_CACHE_DIR/docling")
        pipeline = PdfPipelineOptions(artifacts_path=artifacts, document_timeout=config["parser_timeout"],
                                      do_ocr=True, do_table_structure=True)
        pipeline.ocr_options = TesseractCliOcrOptions(lang=["chi_sim", "eng"])
        options[InputFormat.PDF] = PdfFormatOption(pipeline_options=pipeline)
    converter = DocumentConverter(format_options=options)
    result = converter.convert(path, max_num_pages=config["max_pages"], max_file_size=config["max_upload_bytes"])
    if str(getattr(result.status, "value", result.status)).lower() != "success":
        raise ValueError("Parser did not produce a complete document")
    document = result.document
    if len(document.pages) > config["max_pages"]:
        raise ValueError("Document exceeds configured page limit")
    blocks = []
    worksheet_names = []
    if extension == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet_names = workbook.sheetnames
        workbook.close()
    headings = {}
    total = 0
    for item, depth in document.iterate_items():
        label = str(getattr(item.label, "value", item.label))
        if label in {"page_header", "page_footer"}:
            continue
        if label == "table":
            text = item.export_to_markdown(doc=document)
        else:
            text = getattr(item, "text", "")
        if not text or not text.strip():
            continue
        if label in {"section_header", "title"}:
            level = int(getattr(item, "level", depth + 1))
            headings = {k:v for k,v in headings.items() if k < level}
            headings[level] = text
        heading = " / ".join(headings[k] for k in sorted(headings))
        provenance = [p.model_dump(mode="json") for p in getattr(item, "prov", [])]
        locator = {"format": extension.lstrip("."), "item_ref": getattr(item, "self_ref", ""),
                   "provenance": provenance, "heading_path": heading}
        if provenance:
            locator["page"] = provenance[0].get("page_no")
            if extension == ".pptx":
                locator["slide"] = locator["page"]
            elif extension == ".xlsx" and isinstance(locator["page"], int) and 0 < locator["page"] <= len(worksheet_names):
                locator["sheet"] = worksheet_names[locator["page"] - 1]
                # Docling's Excel backend uses zero-based cell-grid boxes, with exclusive r/b.
                box = provenance[0].get("bbox", {})
                coordinates = [box.get(key) for key in ("l", "t", "r", "b")]
                if label in {"table", "text"} and box.get("coord_origin") == "TOPLEFT" and all(isinstance(value, (int, float)) and float(value).is_integer() for value in coordinates):
                    left, top, right, bottom = map(int, coordinates)
                    if 0 <= left < right and 0 <= top < bottom:
                        from openpyxl.utils.cell import get_column_letter
                        locator["cell_range"] = f"{get_column_letter(left+1)}{top+1}:{get_column_letter(right)}{bottom}"
        # Preserve format-native metadata when the parser exposes it, without inventing coordinates.
        meta = getattr(item, "meta", None)
        if meta is not None and hasattr(meta, "model_dump"):
            locator["native_metadata"] = meta.model_dump(mode="json", exclude_none=True)
        total += len(text)
        if total > config["max_characters"]:
            raise ValueError("Extracted text exceeds configured character limit")
        for offset in range(0, len(text), 12000):
            blocks.append(dict(text=text[offset:offset+12000], heading=heading, kind=label,
                               locator=dict(locator, character_offset=offset)))
    return blocks


if __name__ == "__main__":
    try:
        source, destination, settings_json = sys.argv[1:]
        parsed = _docling_convert(Path(source), json.loads(settings_json))
        Path(destination).write_text(json.dumps(parsed, ensure_ascii=False), encoding="utf-8")
    except Exception as exc:
        print(f"{type(exc).__name__}: {exc}", file=sys.stderr)
        sys.exit(1)
