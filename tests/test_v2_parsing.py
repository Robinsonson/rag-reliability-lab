import io
import subprocess

import pytest

from policy_lab.config import Settings
from policy_lab.parsing import DocumentParser, validate_blocks
from policy_lab.storage import LocalStorage


def settings(tmp_path, **kwargs):
    return Settings(data_dir=tmp_path, model_cache_dir=tmp_path / "models", **kwargs)


def test_markdown_chinese_headings_and_lines(tmp_path):
    path = tmp_path / "policy.md"
    path.write_text("# 差旅制度\n\n住宿上限500元。\n\n## 餐补\n\nMeal allowance 100.", encoding="utf-8")
    blocks = DocumentParser(settings(tmp_path)).parse(path)
    assert [b["ordinal"] for b in blocks] == list(range(4))
    assert blocks[1]["heading"] == "差旅制度"
    assert blocks[3]["heading"] == "差旅制度 / 餐补"
    assert blocks[1]["locator"]["line_start"] == 3
    assert len({b["id"] for b in blocks}) == 4


def test_utf16_and_bounded_long_block(tmp_path):
    path = tmp_path / "policy.txt"
    path.write_text("中文政策" * 4000, encoding="utf-16")
    blocks = DocumentParser(settings(tmp_path)).parse(path)
    assert len(blocks) == 2
    assert max(len(b["text"]) for b in blocks) <= 12000
    assert blocks[1]["locator"]["character_offset"] == 12000


@pytest.mark.parametrize("content,error", [(b"\xff\xffbad", "UTF-8"), (b"abc\x00def", "Binary"), (b"   \n", "readable")])
def test_bad_text_fails_explicitly(tmp_path, content, error):
    path = tmp_path / "bad.txt"
    path.write_bytes(content)
    with pytest.raises(ValueError, match=error):
        DocumentParser(settings(tmp_path)).parse(path)


def test_limits_extensions_and_duplicate_ids(tmp_path):
    path = tmp_path / "large.md"
    path.write_text("a" * 100)
    with pytest.raises(ValueError, match="upload limit"):
        DocumentParser(settings(tmp_path, max_upload_bytes=10)).parse(path)
    with pytest.raises(ValueError, match="character limit"):
        DocumentParser(settings(tmp_path, max_characters=10)).parse(path)
    with pytest.raises(ValueError, match="Unsupported"):
        DocumentParser(settings(tmp_path)).parse(path, "fake.exe")
    with pytest.raises(ValueError, match="Duplicate"):
        validate_blocks([{"id": "same", "text": "a"}, {"id": "same", "text": "b"}], settings(tmp_path))


def test_docling_subprocess_timeout_is_reported(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    path.write_bytes(b"%PDF-1.7")
    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired(args[0], kwargs["timeout"])
    monkeypatch.setattr(subprocess, "run", timeout)
    with pytest.raises(ValueError, match="timeout"):
        DocumentParser(settings(tmp_path)).parse(path)


@pytest.mark.parametrize("key", ["../escape", "/absolute", "C:/absolute", "a\\b"])
def test_local_storage_rejects_path_escape(tmp_path, key):
    storage = LocalStorage(tmp_path / "objects")
    storage.ensure_bucket()
    with pytest.raises(ValueError):
        storage.put_file(key, io.BytesIO(b"secret"))


def test_local_storage_does_not_overwrite_original(tmp_path):
    storage = LocalStorage(tmp_path / "objects")
    storage.ensure_bucket()
    storage.put_file("uploads/id", io.BytesIO(b"original"))
    with pytest.raises(ValueError, match="already exists"):
        storage.put_file("uploads/id", io.BytesIO(b"changed"))
    with storage.open("uploads/id") as source:
        assert source.read() == b"original"


@pytest.mark.parametrize("extension", [".docx", ".pptx", ".xlsx", ".html"])
def test_real_docling_office_html_subprocess(tmp_path, extension):
    pytest.importorskip("docling")
    path = tmp_path / ("policy" + extension)
    expected = "Travel allowance 500 yuan"
    if extension == ".docx":
        from docx import Document
        document = Document()
        document.add_heading("Travel Policy", level=1)
        document.add_paragraph(expected)
        document.save(path)
    elif extension == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = expected
        presentation.save(path)
    elif extension == ".xlsx":
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.title = "Travel"
        workbook.active.append(["Policy", "Amount"])
        workbook.active.append([expected, 500])
        workbook.save(path)
    else:
        path.write_text(f"<!doctype html><html><body><h1>Travel Policy</h1><p>{expected}</p></body></html>", encoding="utf-8")
    blocks = DocumentParser(settings(tmp_path, parser_timeout=90)).parse(path)
    assert expected in "\n".join(block["text"] for block in blocks)
    assert all(block["locator"]["format"] == extension[1:] for block in blocks)
    assert any(block["locator"]["item_ref"] for block in blocks)
    if extension == ".xlsx":
        assert any(block["locator"].get("sheet") == "Travel" and block["locator"].get("cell_range") == "A1:B2" for block in blocks)
    if extension == ".pptx":
        assert any(block["locator"].get("slide") == 1 for block in blocks)


def test_pdf_requires_explicitly_provisioned_models(tmp_path):
    pytest.importorskip("docling")
    path = tmp_path / "scan.pdf"
    path.write_bytes(b"%PDF-1.7\n%%EOF")
    with pytest.raises(ValueError, match="not provisioned"):
        DocumentParser(settings(tmp_path, parser_timeout=90)).parse(path)


def test_expanded_office_archive_is_bounded_before_parser(tmp_path):
    import zipfile
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 10000)
    assert path.stat().st_size < 500
    with pytest.raises(ValueError, match="expanded archive"):
        DocumentParser(settings(tmp_path, max_upload_bytes=500)).parse(path)
